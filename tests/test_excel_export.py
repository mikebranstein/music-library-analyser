"""Tests for Script 10 (Excel workbook export).

Fully offline: builds synthetic Script 06 per-piece records on disk and asserts the flattening of
each sheet's grain (pieces, documents, expected parts, observed parts, action items), the
Collection Summary KPI flattening (including that per-piece arrays are excluded), the empty-input
error, and end-to-end CLI output (workbook + definitions.md written to disk).
"""

from __future__ import annotations

import importlib.util
import json
from pathlib import Path
from typing import Any

import openpyxl
from typer.testing import CliRunner

runner = CliRunner()


def load_module():
    path = Path(__file__).resolve().parent.parent / "scripts" / "10_excel_export.py"
    spec = importlib.util.spec_from_file_location("excel_export_under_test", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


xls = load_module()


# --- Fixtures ------------------------------------------------------------------------------------


def _piece(piece_id: str = "p1", **overrides: Any) -> dict:
    rec = {
        "record_version": "1.1",
        "processing_status": "success",
        "piece_id": piece_id,
        "catalog_number": "001",
        "piece_title_guess": "Test Piece",
        "piece_folder": "001 Test Piece",
        "document_count": 1,
        "ensemble_type": "concert_band",
        "ensemble_display_name": "Concert Band",
        "lookup_status": "matched",
        "identity_match_confidence": 0.9,
        "work_identity": {"resolved": {"composer": "A. Composer", "summary": "A short summary."}},
        "completeness_tier": "complete",
        "completeness_score": 1.0,
        "severity": "ok",
        "needs_review": False,
        "has_score": True,
        "score_expected": True,
        "score_missing": False,
        "expected_part_count": 1,
        "missing_required_count": 0,
        "unexpected_part_count": 0,
        "observed_instrument_count": 1,
        "distinct_instruments": 1,
        "worst_quality_band": "good",
        "quality_summary": {"low_quality_doc_count": 0, "handwritten_doc_count": 0},
        "review_counts": {
            "needs_review_count": 0,
            "low_confidence_count": 0,
            "unmatched_count": 0,
            "duplicate_count": 0,
        },
        "reason_codes": [],
        "recommended_actions": [],
        "sections": ["flutes"],
        "families": ["woodwind"],
        "score_types": ["full"],
        "documents": [],
        "expected_parts": [],
        "observed_parts": [],
        "action_items": [],
    }
    rec.update(overrides)
    return rec


def _write_pieces(tmp_path: Path, records: list[dict]) -> Path:
    piece_dir = tmp_path / "piece_reports"
    piece_dir.mkdir()
    for rec in records:
        (piece_dir / f"{rec['piece_id']}.json").write_text(json.dumps(rec), encoding="utf-8")
    return piece_dir


# --- Row builders ---------------------------------------------------------------------------


def test_build_pieces_rows_flattens_nested_fields():
    rec = _piece()
    rows = xls.build_pieces_rows([rec])
    assert len(rows) == 1
    row = rows[0]
    assert row["piece_id"] == "p1"
    assert row["composer"] == "A. Composer"
    assert row["summary"] == "A short summary."
    assert row["sections"] == "flutes"
    assert row["families"] == "woodwind"
    assert row["score_types"] == "full"


def test_build_documents_rows_one_row_per_document():
    rec = _piece(
        documents=[
            {
                "pdf_path": "001 Test Piece/Flute.pdf",
                "pdf_filename": "Flute.pdf",
                "predicted_part": "Flute",
                "confidence_tier": "high",
                "needs_review": False,
                "is_score": False,
                "score_type": None,
                "duplicate_in_piece": False,
                "page_count": 2,
                "quality_band": "good",
                "quality_score": 95.0,
                "notation_source_type": "printed_original",
                "quality_top_issues": ["low_resolution"],
                "part_sort_key": "1-001",
                "thumbnail_path": "cache/render/x.png",
            }
        ]
    )
    rows = xls.build_documents_rows([rec])
    assert len(rows) == 1
    assert rows[0]["pdf_filename"] == "Flute.pdf"
    assert rows[0]["quality_top_issues"] == "low_resolution"
    assert rows[0]["piece_id"] == "p1"


def test_build_expected_parts_rows_one_row_per_slot():
    rec = _piece(
        expected_parts=[
            {
                "canonical_instrument": "flute",
                "part_index": None,
                "label": "Flute",
                "section": "flutes",
                "required": True,
                "present": False,
                "observed_clefs": [],
                "observed_transpositions": [],
            },
            {
                "canonical_instrument": "oboe",
                "part_index": None,
                "label": "Oboe",
                "section": "double_reeds",
                "required": False,
                "present": True,
                "observed_clefs": ["treble"],
                "observed_transpositions": [],
            },
        ]
    )
    rows = xls.build_expected_parts_rows([rec])
    assert len(rows) == 2
    assert rows[0]["required"] is True
    assert rows[0]["present"] is False
    assert rows[1]["observed_clefs"] == "treble"


def test_build_observed_parts_rows_flattens_instruments():
    rec = _piece(
        observed_parts=[
            {
                "instruments": [{"canonical": "flute", "part_index": None, "section": "flutes"}],
                "clef": None,
                "transposition": None,
                "part_role": "section",
                "predicted_part": "Flute",
                "count": 1,
                "min_confidence": 0.8,
                "max_confidence": 0.8,
                "needs_review": False,
                "duplicate": False,
                "part_sort_key": "1-001",
            }
        ]
    )
    rows = xls.build_observed_parts_rows([rec])
    assert len(rows) == 1
    assert rows[0]["canonical_instruments"] == "flute"
    assert rows[0]["sections"] == "flutes"


def test_build_action_items_rows_one_row_per_target():
    rec = _piece(
        severity="high",
        action_items=[
            {
                "reason_code": "missing_required_parts",
                "action": "Source the missing part(s).",
                "targets": ["Piccolo", "Oboe"],
            }
        ],
    )
    rows = xls.build_action_items_rows([rec])
    assert len(rows) == 2
    assert {r["target"] for r in rows} == {"Piccolo", "Oboe"}
    assert all(r["reason_code"] == "missing_required_parts" for r in rows)


def test_build_action_items_rows_handles_missing_targets():
    rec = _piece(
        action_items=[{"reason_code": "missing_score", "action": "Add a score.", "targets": []}]
    )
    rows = xls.build_action_items_rows([rec])
    assert len(rows) == 1
    assert rows[0]["target"] is None


def test_build_summary_rows_excludes_per_piece_arrays():
    summary = {
        "record_version": "1.2",
        "totals": {"pieces": 5, "pieces_complete": 2},
        "top_missing_sections": [{"section": "percussion", "missing_piece_count": 3}],
        "pieces": [{"piece_id": "p1"}, {"piece_id": "p2"}],
        "attention_pieces": [{"piece_id": "p1"}],
    }
    rows = xls.build_summary_rows(summary)
    metrics = {r["metric"] for r in rows}
    assert "totals.pieces" in metrics
    assert "top_missing_sections[0].section" in metrics
    assert not any(m.startswith("pieces[") for m in metrics)
    assert not any(m.startswith("attention_pieces[") for m in metrics)


def test_build_summary_rows_handles_none():
    assert xls.build_summary_rows(None) == []


# --- Workbook assembly -----------------------------------------------------------------------


def test_build_workbook_has_expected_sheets():
    wb = xls.build_workbook([_piece()], {"totals": {"pieces": 1}})
    assert wb.sheetnames == [
        "Collection Summary",
        "Pieces",
        "Documents",
        "Expected Parts",
        "Observed Parts",
        "Action Items",
    ]
    pieces_ws = wb["Pieces"]
    assert pieces_ws["A1"].value == "piece_id"
    assert pieces_ws.freeze_panes == "A2"
    assert pieces_ws.auto_filter.ref is not None


def test_build_workbook_empty_sheet_placeholder():
    wb = xls.build_workbook([_piece(action_items=[])], None)
    ws = wb["Action Items"]
    assert ws["A1"].value == "(no data)"


# --- CLI --------------------------------------------------------------------------------------


def test_cli_writes_workbook_and_definitions(tmp_path: Path):
    piece_dir = _write_pieces(tmp_path, [_piece("p1"), _piece("p2", catalog_number="002")])
    output_dir = tmp_path / "excel_export"

    result = runner.invoke(
        xls.app,
        [
            "--piece-reports-dir",
            str(piece_dir),
            "--collection-summary",
            str(tmp_path / "does_not_exist.json"),
            "--output-dir",
            str(output_dir),
        ],
    )

    assert result.exit_code == 0, result.output
    workbook_path = output_dir / "music_library_report.xlsx"
    definitions_path = output_dir / "definitions.md"
    assert workbook_path.exists()
    assert definitions_path.exists()
    assert "What Each Column Means" in definitions_path.read_text(encoding="utf-8")

    wb = openpyxl.load_workbook(workbook_path)
    pieces_ws = wb["Pieces"]
    assert pieces_ws.max_row == 3  # header + 2 pieces


def test_cli_errors_when_no_piece_reports(tmp_path: Path):
    empty_dir = tmp_path / "piece_reports"
    empty_dir.mkdir()
    result = runner.invoke(
        xls.app,
        [
            "--piece-reports-dir",
            str(empty_dir),
            "--output-dir",
            str(tmp_path / "excel_export"),
        ],
    )
    assert result.exit_code != 0
