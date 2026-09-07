"""Script 10: Excel workbook export.

Compiles the Script 06 per-piece reports (`data/piece_reports/*.json`, schema 1.1) and the
Script 07 collection summary (`data/collection_reports/summary.json`, schema 1.2) into a single
multi-sheet `.xlsx` workbook so the whole collection can be sorted, filtered, and pivoted in Excel
without touching the underlying JSON/JSONL. Like Scripts 06/07 this is a pure reporting stage: it
never re-derives completeness, quality, severity, or reason codes, it only flattens the nested
per-piece facts into normalized tabular sheets, joined by `piece_id`.

Sheets emitted (one grain per sheet, all joinable on `piece_id`):
  - Collection Summary : headline KPIs from Script 07's summary.json (one column of metrics)
  - Pieces             : one row per piece (completeness, severity, quality, review counts, ...)
  - Documents          : one row per PDF/document (predicted part, quality, notation source, ...)
  - Expected Parts     : one row per expected part slot (required/optional, present/missing)
  - Observed Parts     : one row per detected/observed part group (instrument, confidence, ...)
  - Action Items       : one row per recommended action target (reason code, action, target)

A companion `definitions.md` (column glossary) is written alongside the workbook.
"""

from __future__ import annotations

import json
import logging
import time
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, NamedTuple

import typer
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from scripts._common import (
    atomic_write_text,
    piece_sort_key,
    setup_logging,
)

RECORD_VERSION = "1.0"

# The Script 06 per-piece schema this export is written against. Records at any other version are
# still exported, but their newer/older fields may be missing; main() warns when the input set is
# not uniformly this version (mirrors Scripts 07/08's data-health check).
EXPECTED_PIECE_SCHEMA = "1.1"

app = typer.Typer(add_completion=False)

logger = logging.getLogger("script10.excel_export")

CHECKPOINT_FILENAME = ".piece_report_checkpoint.json"

HEADER_FILL = PatternFill(start_color="FF305496", end_color="FF305496", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")


# --- Input loading ------------------------------------------------------------------------------


class LoadResult(NamedTuple):
    """Outcome of scanning the piece-reports directory."""

    records: list[dict[str, Any]]
    skipped: int


def load_piece_records(piece_reports_dir: Path) -> LoadResult:
    """Read every per-piece JSON record from the Script 06 output directory.

    Only files whose payload is an object with a ``piece_id`` are kept (this skips the checkpoint
    dotfile and any stray files). Unreadable/malformed files and non-piece payloads are skipped and
    counted so the caller can surface a data-health warning. Mirrors Scripts 07/08's loader.
    """
    records: list[dict[str, Any]] = []
    skipped = 0
    if not piece_reports_dir.exists():
        return LoadResult(records, skipped)
    for path in sorted(piece_reports_dir.glob("*.json")):
        if path.name == CHECKPOINT_FILENAME:
            continue
        try:
            with path.open("r", encoding="utf-8") as f:
                payload = json.load(f)
        except (OSError, json.JSONDecodeError) as exc:
            logger.warning("Skipping unreadable piece report %s: %s", path, exc)
            skipped += 1
            continue
        if isinstance(payload, dict) and payload.get("piece_id"):
            records.append(payload)
        else:
            logger.warning("Skipping %s: not a piece report (no piece_id).", path)
            skipped += 1
    return LoadResult(records, skipped)


def load_collection_summary(path: Path) -> dict[str, Any] | None:
    """Read the Script 07 collection summary, if present."""
    if not path.exists():
        return None
    try:
        with path.open("r", encoding="utf-8") as f:
            payload = json.load(f)
    except (OSError, json.JSONDecodeError) as exc:
        logger.warning("Skipping unreadable collection summary %s: %s", path, exc)
        return None
    return payload if isinstance(payload, dict) else None


# --- Row builders (flatten nested Script 06 facts into one grain per sheet) ---------------------


def _sorted_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(records, key=piece_sort_key)


def build_pieces_rows(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """One row per piece: the piece-level facts Script 06 already computed."""
    rows: list[dict[str, Any]] = []
    for rec in _sorted_records(records):
        quality = rec.get("quality_summary") or {}
        review = rec.get("review_counts") or {}
        resolved = ((rec.get("work_identity") or {}).get("resolved")) or {}
        rows.append(
            {
                "piece_id": rec.get("piece_id"),
                "catalog_number": rec.get("catalog_number"),
                "piece_title_guess": rec.get("piece_title_guess"),
                "piece_folder": rec.get("piece_folder"),
                "composer": resolved.get("composer"),
                "ensemble_type": rec.get("ensemble_type"),
                "ensemble_display_name": rec.get("ensemble_display_name"),
                "lookup_status": rec.get("lookup_status"),
                "identity_match_confidence": rec.get("identity_match_confidence"),
                "completeness_tier": rec.get("completeness_tier"),
                "completeness_score": rec.get("completeness_score"),
                "severity": rec.get("severity"),
                "needs_review": bool(rec.get("needs_review")),
                "has_score": bool(rec.get("has_score")),
                "score_expected": bool(rec.get("score_expected")),
                "score_missing": bool(rec.get("score_missing")),
                "document_count": int(rec.get("document_count") or 0),
                "expected_part_count": int(rec.get("expected_part_count") or 0),
                "missing_required_count": int(rec.get("missing_required_count") or 0),
                "unexpected_part_count": int(rec.get("unexpected_part_count") or 0),
                "observed_instrument_count": int(rec.get("observed_instrument_count") or 0),
                "distinct_instruments": int(rec.get("distinct_instruments") or 0),
                "worst_quality_band": rec.get("worst_quality_band"),
                "low_quality_doc_count": int(quality.get("low_quality_doc_count") or 0),
                "handwritten_doc_count": int(quality.get("handwritten_doc_count") or 0),
                "needs_review_count": int(review.get("needs_review_count") or 0),
                "low_confidence_count": int(review.get("low_confidence_count") or 0),
                "unmatched_count": int(review.get("unmatched_count") or 0),
                "duplicate_count": int(review.get("duplicate_count") or 0),
                "reason_codes": "; ".join(rec.get("reason_codes") or []),
                "recommended_actions": " | ".join(rec.get("recommended_actions") or []),
                "sections": "; ".join(rec.get("sections") or []),
                "families": "; ".join(rec.get("families") or []),
                "score_types": "; ".join(rec.get("score_types") or []),
                "summary": resolved.get("summary"),
                "record_version": rec.get("record_version"),
            }
        )
    return rows


def build_documents_rows(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """One row per PDF/document, echoing the piece it belongs to for filtering."""
    rows: list[dict[str, Any]] = []
    for rec in _sorted_records(records):
        piece_id = rec.get("piece_id")
        for doc in rec.get("documents") or []:
            rows.append(
                {
                    "piece_id": piece_id,
                    "catalog_number": rec.get("catalog_number"),
                    "piece_title_guess": rec.get("piece_title_guess"),
                    "pdf_path": doc.get("pdf_path"),
                    "pdf_filename": doc.get("pdf_filename"),
                    "predicted_part": doc.get("predicted_part"),
                    "confidence_tier": doc.get("confidence_tier"),
                    "needs_review": bool(doc.get("needs_review")),
                    "is_score": bool(doc.get("is_score")),
                    "score_type": doc.get("score_type"),
                    "duplicate_in_piece": bool(doc.get("duplicate_in_piece")),
                    "page_count": doc.get("page_count"),
                    "quality_band": doc.get("quality_band"),
                    "quality_score": doc.get("quality_score"),
                    "notation_source_type": doc.get("notation_source_type"),
                    "quality_top_issues": "; ".join(doc.get("quality_top_issues") or []),
                    "part_sort_key": doc.get("part_sort_key"),
                    "thumbnail_path": doc.get("thumbnail_path"),
                }
            )
    return rows


def build_expected_parts_rows(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """One row per expected part slot (required/optional; present or missing)."""
    rows: list[dict[str, Any]] = []
    for rec in _sorted_records(records):
        piece_id = rec.get("piece_id")
        for part in rec.get("expected_parts") or []:
            rows.append(
                {
                    "piece_id": piece_id,
                    "catalog_number": rec.get("catalog_number"),
                    "piece_title_guess": rec.get("piece_title_guess"),
                    "label": part.get("label"),
                    "canonical_instrument": part.get("canonical_instrument"),
                    "section": part.get("section"),
                    "part_index": part.get("part_index"),
                    "required": bool(part.get("required")),
                    "present": bool(part.get("present")),
                    "observed_clefs": "; ".join(part.get("observed_clefs") or []),
                    "observed_transpositions": "; ".join(part.get("observed_transpositions") or []),
                }
            )
    return rows


def build_observed_parts_rows(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """One row per observed/detected part group (as classified by Script 03)."""
    rows: list[dict[str, Any]] = []
    for rec in _sorted_records(records):
        piece_id = rec.get("piece_id")
        for part in rec.get("observed_parts") or []:
            instruments = part.get("instruments") or []
            canonical = "; ".join(i.get("canonical") or "" for i in instruments)
            sections = "; ".join(i.get("section") or "" for i in instruments)
            rows.append(
                {
                    "piece_id": piece_id,
                    "catalog_number": rec.get("catalog_number"),
                    "piece_title_guess": rec.get("piece_title_guess"),
                    "predicted_part": part.get("predicted_part"),
                    "canonical_instruments": canonical,
                    "sections": sections,
                    "clef": part.get("clef"),
                    "transposition": part.get("transposition"),
                    "part_role": part.get("part_role"),
                    "count": part.get("count"),
                    "min_confidence": part.get("min_confidence"),
                    "max_confidence": part.get("max_confidence"),
                    "needs_review": bool(part.get("needs_review")),
                    "duplicate": bool(part.get("duplicate")),
                    "part_sort_key": part.get("part_sort_key"),
                }
            )
    return rows


def build_action_items_rows(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """One row per action-item target: reason code, action text, and the specific target."""
    rows: list[dict[str, Any]] = []
    for rec in _sorted_records(records):
        piece_id = rec.get("piece_id")
        for item in rec.get("action_items") or []:
            targets = item.get("targets") or [None]
            for target in targets:
                rows.append(
                    {
                        "piece_id": piece_id,
                        "catalog_number": rec.get("catalog_number"),
                        "piece_title_guess": rec.get("piece_title_guess"),
                        "severity": rec.get("severity"),
                        "reason_code": item.get("reason_code"),
                        "action": item.get("action"),
                        "target": target,
                    }
                )
    return rows


def build_summary_rows(summary: dict[str, Any] | None) -> list[dict[str, Any]]:
    """Flatten Script 07's nested KPI dict into metric/value rows for a one-glance sheet."""
    rows: list[dict[str, Any]] = []
    if not summary:
        return rows

    def walk(prefix: str, value: Any) -> None:
        if isinstance(value, dict):
            for key, sub in value.items():
                walk(f"{prefix}.{key}" if prefix else key, sub)
        elif isinstance(value, list):
            if value and all(isinstance(v, dict) for v in value):
                for i, item in enumerate(value):
                    walk(f"{prefix}[{i}]", item)
            else:
                rows.append({"metric": prefix, "value": "; ".join(str(v) for v in value)})
        else:
            rows.append({"metric": prefix, "value": value})

    # "pieces" and "attention_pieces" duplicate the per-piece Pieces sheet (one entry per piece
    # needing attention); excluding them keeps this sheet a small, one-glance set of headline KPIs
    # instead of exploding to piece-count rows.
    skip_top_level = {
        "record_version",
        "run_id",
        "processing_status",
        "processing_timestamp",
        "pieces",
        "attention_pieces",
    }
    for key, value in summary.items():
        if key in skip_top_level:
            continue
        walk(key, value)
    return rows


# --- Workbook assembly ---------------------------------------------------------------------------


def _write_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]]) -> Worksheet:
    """Write ``rows`` (a list of flat dicts sharing a column set) as one auto-filtered sheet."""
    ws = wb.create_sheet(title=title[:31])  # Excel sheet-name length limit
    if not rows:
        ws.append(["(no data)"])
        return ws

    columns = list(rows[0].keys())
    ws.append(columns)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        ws.append([row.get(col) for col in columns])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Column widths: cap so a long summary/notes field doesn't blow out the sheet.
    for idx, col in enumerate(columns, start=1):
        max_len = max(
            [len(col)] + [len(str(row.get(col))) for row in rows if row.get(col) is not None]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 10), 60)

    return ws


def build_workbook(
    piece_records: list[dict[str, Any]],
    summary: dict[str, Any] | None,
) -> Workbook:
    """Assemble the full multi-sheet workbook from the loaded Script 06/07 records."""
    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet

    _write_sheet(wb, "Collection Summary", build_summary_rows(summary))
    _write_sheet(wb, "Pieces", build_pieces_rows(piece_records))
    _write_sheet(wb, "Documents", build_documents_rows(piece_records))
    _write_sheet(wb, "Expected Parts", build_expected_parts_rows(piece_records))
    _write_sheet(wb, "Observed Parts", build_observed_parts_rows(piece_records))
    _write_sheet(wb, "Action Items", build_action_items_rows(piece_records))

    return wb


# --- Definitions (Markdown glossary) --------------------------------------------------------------


def render_definitions_md(piece_count: int, document_count: int, generated_at: str) -> str:
    """Render the plain-language definitions/glossary file, stamped with this run's real counts."""
    header = (
        f"This copy of the report covers **{piece_count:,} pieces of music** and "
        f"**{document_count:,} individual sheet-music files**, and was generated on "
        f"**{generated_at}**. These numbers will change the next time the report is refreshed, "
        "as more of the library is reviewed or corrections are made."
    )
    return f"""\
# Music Library Report - What Each Column Means

This document explains the workbook file `music_library_report.xlsx`. It goes with the
spreadsheet and explains, in plain language, what every column on every tab (sheet) means. You do
not need any computer or music-technical background to read this.

{header}

The spreadsheet has six tabs at the bottom: Collection Summary, Pieces, Documents, Expected
Parts, Observed Parts, and Action Items. Every tab except Collection Summary has a column called
`piece_id`. That column is just a way of tying a row on one tab to the same piece of music on
another tab. You do not need to use it yourself, but if you ever want to filter one tab down to a
single piece of music and then find that same piece on another tab, `piece_id` is how you would
match them up.

## How this information was put together, and why it may not be perfect

This report was produced by an automated review of the sheet music library. No person manually
checked every page of every file by hand before this report was created - instead, a computer
process worked through the whole collection on its own, and the results below come from that
process. It is important to understand how that process worked, so you know how much to trust it
and where mistakes are most likely.

For each file, the process looked at the file name and the contents of the pages to work out
which instrument part it was (for example, "1st Clarinet" or "Trombone"). Where a file was a
scanned image rather than typed text, the process used optical character recognition (often
called OCR) - a technology that reads text out of a picture of a page - to figure out what was
written on it. OCR can misread smudged, faint, handwritten, or poor-quality scans, so its results
are not always accurate.

For each piece of music, the process also tried to work out what instrument parts the piece is
supposed to have altogether (for example, whether it should include a part for Oboe, or two
Trumpet parts rather than one). To do this, it searched online for information about the piece
and its publisher, and used an artificial intelligence (AI) system to read and interpret what it
found, in order to identify the piece and decide what its complete, correct set of instrument
parts should be. It also used the AI system to help read and judge the scanned pages themselves,
including checking scan quality and deciding whether a page was printed or handwritten. AI systems
of this kind are generally reliable but are not perfect, and can occasionally be confidently
wrong - for example, matching a piece to the wrong edition, or misjudging a page's quality.

Because of all this, please treat the contents of this report as a well-informed starting point
rather than a guaranteed, error-free record. In particular:

- A piece marked as "missing" a part may already have that part somewhere in the library, if the
  file was misnamed, misread, or not matched correctly.
- A piece marked as "complete" may still be missing something the automated process did not
  catch.
- Composer names, titles, and other identifying details are the process's best judgment and
  should be treated as likely, not certain.
- Scan-quality and handwriting judgments are automated estimates and may not match what a person
  would decide by looking at the page directly.

### How the "correct" list of instrument parts for each piece was worked out

For every piece, the process needed to answer a specific question: what instrument parts should
this piece have, according to its actual published sheet music edition? It tried, in order, a few
different ways of answering that question, and used the first one that gave a confident answer:

1. **Reading the piece's own conductor's score.** If a full or condensed conductor's score was
   found among the piece's own files, the process read the instrument names printed on it
   (typically stacked down the left-hand side of the first page of music) and used that as the
   list of parts the piece should have. This was the most common way an answer was found across
   the collection.
2. **Looking it up on the Wind Repertory Project**, a community-maintained reference website for
   concert band and wind ensemble music, to find the piece's published instrumentation.
3. **Searching more broadly online** for the piece and its publisher, using an AI system to read
   what it found (for example a publisher's product page or catalog listing) and work out the
   correct instrumentation from that.
4. **Reading a picture of the score found online**, when a search turned up an image of the score
   itself rather than a text description - the same picture-reading (OCR) technology described
   above was used on it.

If none of these produced a confident answer, the piece was left without a confirmed "correct"
list of parts, and is flagged for someone to check by hand rather than having a guess recorded.
For this collection, a confirmed answer was found for the large majority of pieces, mostly from
reading the piece's own conductor's score, with the rest split between the Wind Repertory Project,
broader online searches, and pictures of scores found online. Only a smaller portion of pieces
could not be confidently resolved by any of these methods.

Separately, for each individual file (each PDF), the process worked out which single instrument
part that file contains. It started from the file's own name, then checked the printed part name
on the page itself and other clues in the file's text, and used whichever of these it trusted most
for that file. This is how, for example, a file named for one instrument but printed with a
different part name would be caught and corrected.

### Reasons the expected instrumentation list and the actual files may not match

Even when a confident "correct" list of parts was found, it may not perfectly match what is
actually sitting in the library folder for a piece. Common reasons for a mismatch include:

- **The wrong edition was identified.** Many pieces have been published more than once, sometimes
  by different publishers, with different instrumentation. If the process matched a piece to a
  different edition or arrangement than the one actually in the library, the expected part list
  will not match the real files.
- **The library holds a different arrangement than expected.** A folder may contain a school,
  simplified, or custom arrangement that genuinely has a different, smaller, or larger set of
  parts than the standard published edition the process found.
- **A part was found but not recognized as matching.** A file may already be the missing part, but
  named or labeled in a way the process did not recognize as matching the expected instrument
  name, so it was counted as missing when it is not.
- **A part exists but is filed under the wrong piece.** Sheet music is sometimes misfiled into the
  wrong folder; a part could be missing from where it is expected but sitting in another piece's
  folder instead.
- **The score itself lists optional, cued, or alternate parts.** Some instrument names on a score
  represent optional extras or alternate versions of another part (for example, a part that can be
  covered by a cue in another instrument's part). Depending on how clearly this was printed, the
  process may or may not have judged such a part as strictly required.
- **The reading of the score or the online source was itself wrong.** Because this all depends on
  optical character recognition and an AI system reading pages, catalog listings, or images, any
  step in the process can misread or misinterpret what it was given, producing an expected part
  list that does not reflect the real published edition at all.

If anything in this report looks surprising or does not match what you know about the collection,
it is worth checking the actual file(s) in question before relying on the report alone.

## Tab: Collection Summary

This tab has two columns only: `metric` (a short name describing what is being counted) and
`value` (the number or answer for it). It gives you the big-picture totals for the whole
collection, for example how many pieces of music there are in total, how many are missing parts,
and which instrument sections are most often missing across the whole library. For the detail
behind any of these totals, look at the Pieces tab instead.

| Column | What it means |
| --- | --- |
| `metric` | A short label describing what this row is counting. |
| `value` | The number or answer for that count. |

## Tab: Pieces

This is the main tab. Each row is one piece of music (one folder of sheet music files).

| Column | What it means |
| --- | --- |
| `piece_id` | An internal reference code for this piece of music. Used only to match rows between tabs. |
| `catalog_number` | The library's catalog number for this piece, taken from its folder name. |
| `piece_title_guess` | The title of the piece, as best determined by the review process. |
| `piece_folder` | The name of the folder this piece's files are stored in. |
| `composer` | The composer's name, if it could be confidently identified. |
| `ensemble_type` / `ensemble_display_name` | What kind of musical group this piece is written for (for example, Concert Band). |
| `lookup_status` | Whether a match for this piece was found when checking outside sources (publisher listings, catalogs, etc.) to confirm what instrument parts it should have. One of: `matched` (a confident match was found), `low_confidence` (a possible match was found but is not certain), `no_match` (no match could be found), `disabled` (this check was turned off for this run), or `error` (the check could not be completed). |
| `identity_match_confidence` | How confident the process is that it found the right piece of music when checking outside sources, shown as a number between 0 (not confident at all) and 1 (fully confident). |
| `completeness_tier` | An overall label for how complete this piece's set of instrument parts is: `complete`, `near_complete`, `incomplete`, `severely_incomplete`, or `unknown` (not enough information to judge). |
| `completeness_score` | The share of the required instrument parts that are actually present, shown as a number between 0 (none present) and 1 (all present). |
| `severity` | A simple flag for how much attention this piece needs: `ok` (no concerns), `review` (worth a look), or `high` (needs attention soon). |
| `needs_review` | Marked True if something about this piece needs a person to check it by hand. |
| `has_score` | Marked True if at least one full conductor's score was found for this piece. |
| `score_expected` | Marked True if a conductor's score is expected to exist for this piece. |
| `score_missing` | Marked True if a conductor's score was expected but could not be found. |
| `document_count` | The number of separate sheet-music files found for this piece. |
| `expected_part_count` | The number of individual instrument parts this piece is expected to have, based on the published edition. |
| `missing_required_count` | The number of required instrument parts that could not be found among this piece's files. |
| `unexpected_part_count` | The number of instrument parts found that do not appear to belong to this piece's expected set (for example, an extra or mismatched part). |
| `observed_instrument_count` | The number of different instruments actually found among this piece's files. |
| `distinct_instruments` | The number of different instruments involved overall, counting both what is expected and what was found. |
| `worst_quality_band` | The scan quality of the worst file found for this piece: `good`, `fair`, `poor`, or `unknown` (quality could not be judged). |
| `low_quality_doc_count` | The number of files for this piece with poor or fair scan quality. |
| `handwritten_doc_count` | The number of files for this piece that appear to be handwritten rather than printed. |
| `needs_review_count` | The number of this piece's files flagged as needing a person to check them by hand. |
| `low_confidence_count` | The number of this piece's files where the instrument name could not be identified with confidence. |
| `unmatched_count` | The number of this piece's files where no instrument name could be matched at all. |
| `duplicate_count` | The number of this piece's files that appear to be duplicates of another file for the same instrument part. |
| `reason_codes` | A list of the specific reasons this piece needs attention. See "Reasons for attention" below for what each one means. |
| `recommended_actions` | Plain-language suggestions for what to do about each reason this piece needs attention. |
| `sections` | The instrument sections involved in this piece (for example, clarinets, low brass, percussion). |
| `families` | The broader instrument families involved (for example, woodwind, brass, percussion). |
| `score_types` | The types of conductor's scores found for this piece (for example, full score, condensed score). |
| `summary` | A short two-to-three sentence description of the piece, when one could be found. |
| `record_version` | An internal version number for how this row was produced. Not meaningful day-to-day; only useful if something looks inconsistent and needs investigating. |

### Reasons for attention

These are the specific codes that can appear in the `reason_codes` column above.

| Code | What it means |
| --- | --- |
| `missing_score` | A conductor's score was expected for this piece, but none could be found. |
| `missing_required_parts` | One or more instrument parts that this piece needs are missing. |
| `unexpected_parts` | One or more parts were found that do not seem to match what this piece's edition should have. |
| `low_confidence_parts` | One or more parts have instrument names that could not be identified with confidence. |
| `duplicate_parts` | The same instrument part appears more than once, which may mean a duplicate or redundant copy. |
| `instrumentation_unresolved` | It was not possible to confirm, from any outside source, what instrument parts this piece is supposed to have. |

## Tab: Documents

Each row is one sheet-music file (one PDF). Use this tab to look at individual files - for
example, to find files with poor scan quality or to see which instrument each file is for.

| Column | What it means |
| --- | --- |
| `piece_id` / `catalog_number` / `piece_title_guess` | Identify which piece of music this file belongs to. |
| `pdf_path` | Where this file is stored, relative to the main library folder. |
| `pdf_filename` | The file's name only. |
| `predicted_part` | The instrument part this file is believed to be, based on its file name and its contents. |
| `confidence_tier` | How confident that identification is: `high`, `medium`, or `low`. |
| `needs_review` | Marked True if this file needs a person to check it by hand. |
| `is_score` | Marked True if this file is a conductor's score rather than a single instrument part. |
| `score_type` | If this file is a score, what kind: `full`, `condensed`, or `conductor`. Left blank if it is not a score. |
| `duplicate_in_piece` | Marked True if this file appears to be a duplicate of another file for the same piece. |
| `page_count` | The number of pages in this file. |
| `quality_band` | An overall scan-quality rating for this file: `good`, `fair`, `poor`, or `unknown`. |
| `quality_score` | A scan-quality score for this file, from 0 (worst) to 100 (best). |
| `notation_source_type` | Whether the music notation in this file appears to be `printed_original`, `handwritten`, or `mixed_or_uncertain`. |
| `quality_top_issues` | A list of the most significant quality problems found in this file, if any (for example, poor resolution or heavy blurriness). |
| `part_sort_key` | An internal ordering value used to sort parts in a conventional score order (woodwinds, then brass, then percussion, and so on). Not meaningful to read directly. |
| `thumbnail_path` | Where a small preview picture of this file's first page is stored, for a quick visual check. |

## Tab: Expected Parts

Each row is one instrument part that a piece of music is expected to have, based on its published
edition, whether or not that part was actually found. To see every required part that is missing
across the whole library, filter this tab so `present` is FALSE and `required` is TRUE.

| Column | What it means |
| --- | --- |
| `piece_id` / `catalog_number` / `piece_title_guess` | Identify which piece of music this expected part belongs to. |
| `label` | The name of the part as it would appear in the published edition (for example, "2nd Clarinet"). |
| `canonical_instrument` | The standardized instrument name (for example, "clarinet"). |
| `section` | Which instrument section this part belongs to (for example, clarinets, low brass, percussion). |
| `part_index` | The part's number within its section, if it has one (for example, "2" for "2nd Clarinet"). |
| `required` | Marked True if this part is required by the edition, rather than optional. |
| `present` | Marked True if a matching file was actually found for this part. |
| `observed_clefs` | The musical clef(s) seen in the matching file, if the part was found. |
| `observed_transpositions` | The instrument key/transposition(s) seen in the matching file, if the part was found. |

## Tab: Observed Parts

Each row is one group of instrument parts that was actually found among a piece's files, whether
or not it matches an expected part. Use this tab to see exactly what was found, including any
extra or duplicate parts.

| Column | What it means |
| --- | --- |
| `piece_id` / `catalog_number` / `piece_title_guess` | Identify which piece of music this observed part belongs to. |
| `predicted_part` | The instrument part name as identified from the file (for example, "1st Bb Clarinet"). |
| `canonical_instruments` | The standardized instrument name(s) matched to this part. |
| `sections` | The instrument section(s) this part belongs to. |
| `clef` | The musical clef detected for this part, if any. |
| `transposition` | The instrument key/transposition detected for this part, if any. |
| `part_role` | The part's role, for example a standard section part, a solo part, or a divided (split) part. |
| `count` | How many files were grouped together under this part. |
| `min_confidence` / `max_confidence` | The lowest and highest confidence, from 0 to 1, in identifying this part correctly across its files. |
| `needs_review` | Marked True if any file in this group needs a person to check it by hand. |
| `duplicate` | Marked True if this part appears more than once within the same piece. |
| `part_sort_key` | An internal ordering value used to sort parts in a conventional score order. Not meaningful to read directly. |

## Tab: Action Items

Each row is one specific, concrete follow-up task, such as "find this missing part" or "confirm
this extra part is correct." This is the best tab to work from if you want a simple task list.
You can filter it by `reason_code` or sort it by `severity`.

| Column | What it means |
| --- | --- |
| `piece_id` / `catalog_number` / `piece_title_guess` | Identify which piece of music this task belongs to. |
| `severity` | How much attention the overall piece needs: `ok`, `review`, or `high`. |
| `reason_code` | Which issue this task addresses. See "Reasons for attention" above for what each code means. |
| `action` | A plain-language description of what to do. |
| `target` | The specific file or part this task concerns. May be blank if the task applies to the whole piece rather than one file. |

## Tips for using this workbook in Excel

You do not need any special skills to explore this report - a few basic Excel features go a long
way:

- **Filtering.** Every tab has small drop-down arrows on its top row of column headings. Click one
  to show only the rows matching a value you pick (for example, on the Pieces tab, filter
  `severity` to `high` to see only the pieces needing the most attention).
- **Sorting.** Click a column heading's drop-down and choose to sort largest-to-smallest or
  smallest-to-largest (for example, sort the Pieces tab by `missing_required_count` to see the
  pieces missing the most parts at the top).
- **Finding one piece everywhere.** Filter or search for a piece's `catalog_number` or
  `piece_title_guess` on any tab to see everything the report recorded about that one piece,
  including its individual files (Documents tab), its expected parts (Expected Parts tab), and any
  follow-up tasks (Action Items tab).
- **Counting or totaling a filtered set.** Select a range of cells in a numeric column (for
  example `missing_required_count`) and Excel shows a quick sum, average, and count for the
  selected cells at the bottom-right of the window - useful for quickly sizing up a filtered group
  of pieces.
- The first row of each tab is frozen (it stays visible while you scroll), so column headings are
  always in view.

## A few music terms explained

A handful of terms come up repeatedly in this report. In case any are unfamiliar:

- **Part** - the sheet music written for one specific instrument (for example, the Clarinet part
  contains only the notes the clarinet player plays, not the whole piece).
- **Score** - a single document that shows every instrument's part together, all lined up, usually
  used by the conductor rather than an individual player. A "full score" shows every part
  separately; a "condensed score" combines several parts onto fewer staff lines to save space.
- **Section** - a group of similar instruments, such as "clarinets," "low brass," or "percussion."
  A piece's instrumentation is often organized section by section.
- **Instrumentation** - the complete list of which instruments (and how many of each) a piece of
  music is written for.
- **Transposition** - some instruments (for example, the Bb Clarinet or the F Horn) read and play
  music that is written in a different key than it actually sounds. "Transposition" refers to
  which such instrument a part is written for.
- **Clef** - the symbol at the start of a staff (the five lines music is written on) that tells a
  player how to read the pitches; different instruments conventionally use different clefs (most
  commonly treble or bass clef).
- **Edition / arrangement** - the same piece of music can be published more than once, by
  different publishers or arrangers, with different instrumentation, difficulty, or layout. This
  report is only meaningful in relation to whichever specific edition it identified for a piece.
"""


# --- CLI ------------------------------------------------------------------------------------------


@app.command()
def main(
    piece_reports_dir: Path = typer.Option(
        Path("data/piece_reports"), help="Directory of Script 06 per-piece JSON records"
    ),
    collection_summary: Path = typer.Option(
        Path("data/collection_reports/summary.json"), help="Script 07 collection summary JSON"
    ),
    output_dir: Path = typer.Option(
        Path("outputs/excel_export"), help="Directory for the workbook and definitions.md"
    ),
    workbook_name: str = typer.Option(
        "music_library_report.xlsx", help="Output workbook file name"
    ),
    log_level: str = typer.Option("INFO", help="DEBUG, INFO, WARNING, ERROR"),
) -> None:
    """Compile the pipeline's per-piece reports into one multi-sheet Excel workbook."""
    setup_logging(log_level)
    start_time = time.perf_counter()
    run_id = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")

    piece_reports_dir = piece_reports_dir.resolve()
    collection_summary = collection_summary.resolve()
    output_dir = output_dir.resolve()

    loaded = load_piece_records(piece_reports_dir)
    records = loaded.records
    if not records:
        raise typer.BadParameter(
            f"No piece reports found in {piece_reports_dir}. Run Script 06 first."
        )

    logger.info("Compiling %d piece report(s) from %s.", len(records), piece_reports_dir)

    if loaded.skipped:
        logger.warning(
            "Skipped %d unreadable/invalid file(s) in %s.", loaded.skipped, piece_reports_dir
        )
    versions = {rec.get("record_version") or "unknown" for rec in records}
    unexpected_versions = versions - {EXPECTED_PIECE_SCHEMA}
    if unexpected_versions:
        logger.warning(
            "Piece reports include schema version(s) other than %s (%s); some columns may be "
            "missing for those rows. Re-run Script 06 to refresh them.",
            EXPECTED_PIECE_SCHEMA,
            sorted(unexpected_versions),
        )

    summary = load_collection_summary(collection_summary)
    if summary is None:
        logger.warning(
            "No collection summary found at %s; the Collection Summary sheet will be empty. "
            "Run Script 07 first for a populated summary.",
            collection_summary,
        )

    wb = build_workbook(records, summary)

    document_count = sum(int(rec.get("document_count") or 0) for rec in records)
    generated_at = datetime.now(UTC).strftime("%B %d, %Y")

    output_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = output_dir / workbook_name
    wb.save(workbook_path)
    atomic_write_text(
        output_dir / "definitions.md",
        render_definitions_md(len(records), document_count, generated_at),
    )

    logger.info(
        "Excel export completed: run_id=%s pieces=%d elapsed=%.1fs workbook=%s",
        run_id,
        len(records),
        time.perf_counter() - start_time,
        workbook_path,
    )


if __name__ == "__main__":
    app()
